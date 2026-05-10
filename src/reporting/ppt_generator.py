"""
PPT 报告生成器
支持两种模板：技术快讯（One-Pager）和深度洞察（Deep Dive）
"""
import os
from datetime import datetime
from pathlib import Path
from typing import Dict, List

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

from src.utils.config import config

# 配色方案（科技蓝风格）
COLOR_PRIMARY = RGBColor(0x1A, 0x23, 0x7E)      # 深蓝
COLOR_ACCENT = RGBColor(0x00, 0x96, 0xC7)       # 亮蓝
COLOR_TEXT = RGBColor(0x33, 0x33, 0x33)         # 深灰
COLOR_LIGHT_BG = RGBColor(0xF0, 0xF4, 0xF8)    # 浅蓝背景
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)

class PPTGenerator:
    def __init__(self):
        self.output_dir = Path(config.reporting.get("output_dir", "./data/reports"))
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
    
    def generate_flash_brief(self, content: Dict, docs: List[Dict]) -> str:
        """生成技术快讯（单页）"""
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        
        slide_layout = self.prs.slide_layouts[6]  # 空白布局
        slide = self.prs.slides.add_slide(slide_layout)
        
        # 标题栏背景
        title_box = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.2)
        )
        title_box.fill.solid()
        title_box.fill.fore_color.rgb = COLOR_PRIMARY
        title_box.line.fill.background()
        
        # 标题
        title_tf = title_box.text_frame
        title_tf.text = content.get("技术名称", "技术快讯")
        title_tf.paragraphs[0].font.size = Pt(32)
        title_tf.paragraphs[0].font.bold = True
        title_tf.paragraphs[0].font.color.rgb = COLOR_WHITE
        title_tf.paragraphs[0].alignment = PP_ALIGN.LEFT
        title_tf.margin_left = Inches(0.5)
        title_tf.margin_top = Inches(0.3)
        
        # 日期标签
        date_label = slide.shapes.add_textbox(Inches(10.5), Inches(0.35), Inches(2.5), Inches(0.5))
        date_tf = date_label.text_frame
        date_tf.text = datetime.now().strftime("%Y-%m-%d")
        date_tf.paragraphs[0].font.size = Pt(14)
        date_tf.paragraphs[0].font.color.rgb = COLOR_WHITE
        date_tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
        
        # 关键创新
        y_pos = 1.5
        self._add_section_title(slide, "关键创新", Inches(0.5), Inches(y_pos))
        y_pos += 0.5
        innovations = content.get("关键创新", [])
        for item in innovations:
            self._add_bullet_text(slide, f"• {item}", Inches(0.7), Inches(y_pos), width=Inches(12))
            y_pos += 0.5
        
        # 成熟度评估
        y_pos += 0.2
        self._add_section_title(slide, "成熟度评估", Inches(0.5), Inches(y_pos))
        y_pos += 0.5
        trl = content.get("成熟度评估", {})
        trl_text = f"TRL-{trl.get('等级', 'N/A')} | {trl.get('理由', '')}"
        self._add_bullet_text(slide, trl_text, Inches(0.7), Inches(y_pos), width=Inches(12))
        
        # 主要玩家
        y_pos += 0.7
        self._add_section_title(slide, "主要玩家", Inches(0.5), Inches(y_pos))
        y_pos += 0.5
        players = content.get("主要玩家", [])
        self._add_bullet_text(slide, ", ".join(players), Inches(0.7), Inches(y_pos), width=Inches(12))
        
        # 行动建议（底部高亮框）
        y_pos += 0.8
        advice_box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(y_pos), Inches(12.3), Inches(1.0)
        )
        advice_box.fill.solid()
        advice_box.fill.fore_color.rgb = COLOR_LIGHT_BG
        advice_box.line.color.rgb = COLOR_ACCENT
        advice_box.line.width = Pt(2)
        
        advice_tf = advice_box.text_frame
        advice_tf.text = f"💡 建议：{content.get('一句话建议', '建议持续关注该技术方向')}"
        advice_tf.paragraphs[0].font.size = Pt(18)
        advice_tf.paragraphs[0].font.bold = True
        advice_tf.paragraphs[0].font.color.rgb = COLOR_PRIMARY
        advice_tf.paragraphs[0].alignment = PP_ALIGN.LEFT
        advice_tf.word_wrap = True
        advice_tf.margin_left = Inches(0.3)
        advice_tf.margin_top = Inches(0.2)
        
        # 来源
        y_pos += 1.3
        sources = [f"{d.get('source', '')}: {d.get('title', '')[:40]}..." for d in docs[:3]]
        self._add_bullet_text(slide, "来源：" + "; ".join(sources), Inches(0.5), Inches(y_pos), 
                              width=Inches(12), font_size=10, color=RGBColor(0x88, 0x88, 0x88))
        
        filename = f"FlashBrief_{datetime.now().strftime('%Y%m%d')}_{self._safe_name(content.get('技术名称', 'unknown'))}.pptx"
        filepath = self.output_dir / filename
        self.prs.save(filepath)
        return str(filepath)
    
    def generate_deep_dive(self, content: Dict, docs: List[Dict]) -> str:
        """生成深度洞察报告（多页）"""
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        
        # 1. 封面
        self._create_cover(content)
        
        # 2. 背景与动机
        self._create_content_slide("背景与动机", content.get("背景与动机", ""))
        
        # 3. 技术原理
        self._create_content_slide("技术原理", content.get("技术原理", ""))
        
        # 4. 业界进展
        milestones = content.get("业界进展", [])
        self._create_list_slide("业界进展时间线", milestones)
        
        # 5. 竞争格局
        matrix = content.get("竞争格局矩阵", [])
        self._create_table_slide("竞争格局", matrix)
        
        # 6. 优劣势分析
        swot = content.get("优劣势分析", {})
        self._create_swot_slide(swot)
        
        # 7. 启示
        implications = content.get("对芯片设计的启示", [])
        self._create_list_slide("对芯片设计的启示", implications)
        
        # 8. 行动建议
        actions = content.get("行动建议", {})
        self._create_actions_slide(actions)
        
        filename = f"DeepDive_{datetime.now().strftime('%Y%m%d')}_{self._safe_name(content.get('封面标题', 'unknown'))}.pptx"
        filepath = self.output_dir / filename
        self.prs.save(filepath)
        return str(filepath)
    
    def _create_cover(self, content: Dict):
        slide_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(slide_layout)
        
        # 背景
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_PRIMARY
        bg.line.fill.background()
        
        # 标题
        title = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.333), Inches(1.5))
        tf = title.text_frame
        tf.text = content.get("封面标题", "深度洞察报告")
        tf.paragraphs[0].font.size = Pt(44)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = COLOR_WHITE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        
        # 核心结论
        conclusion = slide.shapes.add_textbox(Inches(1), Inches(4.2), Inches(11.333), Inches(1.0))
        ctf = conclusion.text_frame
        ctf.text = content.get("核心结论", "")
        ctf.paragraphs[0].font.size = Pt(20)
        ctf.paragraphs[0].font.color.rgb = RGBColor(0xCC, 0xDD, 0xFF)
        ctf.paragraphs[0].alignment = PP_ALIGN.CENTER
        
        # 日期
        date = slide.shapes.add_textbox(Inches(1), Inches(6.5), Inches(11.333), Inches(0.5))
        dtf = date.text_frame
        dtf.text = datetime.now().strftime("%Y年%m月%d日")
        dtf.paragraphs[0].font.size = Pt(16)
        dtf.paragraphs[0].font.color.rgb = COLOR_WHITE
        dtf.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    def _create_content_slide(self, title: str, text: str):
        slide_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(slide_layout)
        
        self._add_section_title(slide, title, Inches(0.5), Inches(0.4))
        
        content = slide.shapes.add_textbox(Inches(0.7), Inches(1.2), Inches(12), Inches(5.8))
        tf = content.text_frame
        tf.text = text
        tf.paragraphs[0].font.size = Pt(18)
        tf.paragraphs[0].font.color.rgb = COLOR_TEXT
        tf.word_wrap = True
    
    def _create_list_slide(self, title: str, items: List[str]):
        slide_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(slide_layout)
        
        self._add_section_title(slide, title, Inches(0.5), Inches(0.4))
        
        y = 1.2
        for item in items:
            self._add_bullet_text(slide, f"• {item}", Inches(0.7), Inches(y), width=Inches(12))
            y += 0.6
    
    def _create_table_slide(self, title: str, rows: List[Dict]):
        slide_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(slide_layout)
        
        self._add_section_title(slide, title, Inches(0.5), Inches(0.4))
        
        if not rows:
            return
        
        headers = list(rows[0].keys())
        table = slide.shapes.add_table(len(rows)+1, len(headers), Inches(0.7), Inches(1.2), Inches(12), Inches(0.6)).table
        
        # 表头
        for i, h in enumerate(headers):
            cell = table.cell(0, i)
            cell.text = h
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_PRIMARY
            paragraph = cell.text_frame.paragraphs[0]
            paragraph.font.size = Pt(14)
            paragraph.font.bold = True
            paragraph.font.color.rgb = COLOR_WHITE
        
        # 数据
        for r_idx, row in enumerate(rows, 1):
            for c_idx, h in enumerate(headers):
                cell = table.cell(r_idx, c_idx)
                cell.text = str(row.get(h, ""))
                paragraph = cell.text_frame.paragraphs[0]
                paragraph.font.size = Pt(12)
                paragraph.font.color.rgb = COLOR_TEXT
                if r_idx % 2 == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = COLOR_LIGHT_BG
    
    def _create_swot_slide(self, swot: Dict):
        slide_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(slide_layout)
        
        self._add_section_title(slide, "优劣势与风险分析", Inches(0.5), Inches(0.4))
        
        items = [
            ("优势 (Strengths)", swot.get("优势", []), COLOR_ACCENT),
            ("劣势 (Weaknesses)", swot.get("劣势", []), RGBColor(0xE0, 0x6F, 0x1F)),
            ("机会 (Opportunities)", swot.get("机会", []), RGBColor(0x2E, 0x7D, 0x32)),
            ("威胁 (Threats)", swot.get("威胁", []), RGBColor(0xC6, 0x28, 0x28)),
        ]
        
        positions = [(0.5, 1.2), (6.8, 1.2), (0.5, 4.0), (6.8, 4.0)]
        
        for (label, vals, color), (x, y) in zip(items, positions):
            box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(6), Inches(2.5))
            box.fill.solid()
            box.fill.fore_color.rgb = COLOR_LIGHT_BG
            box.line.color.rgb = color
            box.line.width = Pt(2)
            
            tf = box.text_frame
            tf.text = label
            tf.paragraphs[0].font.size = Pt(16)
            tf.paragraphs[0].font.bold = True
            tf.paragraphs[0].font.color.rgb = color
            
            for v in vals:
                p = tf.add_paragraph()
                p.text = f"• {v}"
                p.font.size = Pt(12)
                p.font.color.rgb = COLOR_TEXT
                p.space_after = Pt(4)
    
    def _create_actions_slide(self, actions: Dict):
        slide_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(slide_layout)
        
        self._add_section_title(slide, "行动建议", Inches(0.5), Inches(0.4))
        
        phases = [
            ("短期（3个月）", actions.get("短期", ""), RGBColor(0x2E, 0x7D, 0x32)),
            ("中期（1年）", actions.get("中期", ""), RGBColor(0xF9, 0xA8, 0x25)),
            ("长期（3年）", actions.get("长期", ""), COLOR_PRIMARY),
        ]
        
        y = 1.2
        for label, text, color in phases:
            box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(y), Inches(12), Inches(1.6))
            box.fill.solid()
            box.fill.fore_color.rgb = COLOR_LIGHT_BG
            box.line.color.rgb = color
            box.line.width = Pt(2)
            
            tf = box.text_frame
            tf.text = f"{label}\n{text}"
            tf.paragraphs[0].font.size = Pt(16)
            tf.paragraphs[0].font.bold = True
            tf.paragraphs[0].font.color.rgb = color
            tf.paragraphs[1].font.size = Pt(14)
            tf.paragraphs[1].font.color.rgb = COLOR_TEXT
            tf.word_wrap = True
            tf.margin_left = Inches(0.2)
            tf.margin_top = Inches(0.1)
            
            y += 2.0
    
    def _add_section_title(self, slide, text: str, left, top):
        box = slide.shapes.add_textbox(left, top, Inches(12), Inches(0.5))
        tf = box.text_frame
        tf.text = text
        tf.paragraphs[0].font.size = Pt(22)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = COLOR_PRIMARY
    
    def _add_bullet_text(self, slide, text: str, left, top, width=Inches(12), font_size=14, color=COLOR_TEXT):
        box = slide.shapes.add_textbox(left, top, width, Inches(0.5))
        tf = box.text_frame
        tf.text = text
        tf.paragraphs[0].font.size = Pt(font_size)
        tf.paragraphs[0].font.color.rgb = color
        tf.word_wrap = True
    
    def _safe_name(self, name: str) -> str:
        return "".join(c if c.isalnum() or c in "-_" else "_" for c in name)[:50]
